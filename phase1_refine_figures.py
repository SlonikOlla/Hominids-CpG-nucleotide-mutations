#!/usr/bin/env python3

import csv
from pathlib import Path

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from pptx import Presentation
from pptx.util import Inches
from PIL import Image

ROOT = Path("results/phase1_genome/aggregated")
OUT = ROOT / "figures"
OUT.mkdir(parents=True, exist_ok=True)

LINEAGES = ["human", "chimp", "bonobo", "gorilla"]
LABELS = ["Human", "Chimpanzee", "Bonobo", "Gorilla"]
CONTEXTS = ["CG", "CHG", "CHH"]

COLORS = {
    "human": "#B44E49",
    "chimp": "#427EAD",
    "bonobo": "#549B73",
    "gorilla": "#8A67A5",
}

plt.rcParams.update({
    "font.size": 10,
    "axes.spines.top": False,
    "axes.spines.right": False,
    "pdf.fonttype": 42,
    "figure.facecolor": "white",
    "savefig.facecolor": "white",
})

def read_tsv(path):
    with open(path, newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))

def save(fig, name):
    fig.savefig(OUT / (name + ".png"), dpi=300,
                bbox_inches="tight")
    fig.savefig(OUT / (name + ".pdf"),
                bbox_inches="tight")
    plt.close(fig)
    print("CREATED:", name)

# Read exact genome-wide 5-mer data.
rows = read_tsv(ROOT / "genome_exact_5mer_summary.tsv")

# Check available columns.
print("INPUT COLUMNS:", list(rows[0].keys()))

# Select posterior probability >= 0.95.
rows = [
    r for r in rows
    if abs(float(r["posterior_threshold"]) - 0.95) < 1e-8
]

def get_rate(row):
    return (
        float(row["C>T"])
        / float(row["opportunities"])
        * 1_000_000
    )

# Keep motifs with valid rates in all four lineages.
data = {}

for context in CONTEXTS:
    data[context] = {}

    for r in rows:
        if r["context"] != context:
            continue

        motif = r["ancestral_5mer"]
        lineage = r["lineage"]

        if lineage not in LINEAGES:
            continue

        if float(r["opportunities"]) <= 0:
            continue

        data[context].setdefault(motif, {})[lineage] = get_rate(r)

    data[context] = {
        motif: values
        for motif, values in data[context].items()
        if all(lineage in values for lineage in LINEAGES)
    }

# Select top and bottom motifs by mean within-lineage rank.
selected = {}

for context in CONTEXTS:
    motifs = list(data[context])

    ranks = {}

    for lineage in LINEAGES:
        ordered = sorted(
            motifs,
            key=lambda m: data[context][m][lineage]
        )
        ranks[lineage] = {
            motif: i for i, motif in enumerate(ordered)
        }

    mean_rank = {
        motif: np.mean([
            ranks[lineage][motif] for lineage in LINEAGES
        ])
        for motif in motifs
    }

    low = sorted(motifs, key=lambda m: mean_rank[m])[:3]
    high = sorted(motifs, key=lambda m: mean_rank[m])[-3:]

    selected[context] = low + high

    print(context, "LOW:", low, "HIGH:", high)

# Export exact values.
table_path = OUT / "Figure4_selected_motifs.tsv"

with open(table_path, "w", newline="") as f:
    writer = csv.writer(f, delimiter="\t")
    writer.writerow([
        "context", "motif", "lineage",
        "CT_rate_per_million"
    ])

    for context in CONTEXTS:
        for motif in selected[context]:
            for lineage in LINEAGES:
                writer.writerow([
                    context, motif, lineage,
                    f"{data[context][motif][lineage]:.6f}"
                ])

print("CREATED:", table_path)

# Create revised Figure 4.
fig = plt.figure(figsize=(15, 11))
gs = fig.add_gridspec(
    3, 2,
    width_ratios=[1.5, 1],
    hspace=0.42,
    wspace=0.23
)

for i, context in enumerate(CONTEXTS):

    motifs = sorted(data[context])

    matrix = np.array([
        [data[context][motif][lineage] for motif in motifs]
        for lineage in LINEAGES
    ])

    # Normalize each lineage to its context-specific mean.
    normalized = matrix / matrix.mean(axis=1, keepdims=True)

    ax = fig.add_subplot(gs[i, 0])

    im = ax.imshow(
        normalized,
        aspect="auto",
        cmap="viridis",
        vmin=0.4,
        vmax=1.8
    )

    ax.set_yticks(range(4))
    ax.set_yticklabels(LABELS)
    ax.set_xticks([])

    ax.set_title(
        f"{context}: {len(motifs)} ancestral 5-mers",
        fontsize=12
    )

    fig.colorbar(
        im,
        ax=ax,
        fraction=0.025,
        pad=0.02,
        label="Rate / lineage-context mean"
    )

    ax2 = fig.add_subplot(gs[i, 1])

    subset = selected[context]
    y = np.arange(len(subset))
    height = 0.18

    for j, lineage in enumerate(LINEAGES):
        values = [
            data[context][motif][lineage]
            for motif in subset
        ]

        ax2.barh(
            y + (j - 1.5) * height,
            values,
            height,
            color=COLORS[lineage],
            label=LABELS[j]
        )

    ax2.set_yticks(y)
    ax2.set_yticklabels(subset)
    ax2.invert_yaxis()
    ax2.set_xlabel("C→T substitutions per million opportunities")
    ax2.set_title("Selected low- and high-rate motifs")
    ax2.grid(axis="x", alpha=0.2)
    ax2.set_axisbelow(True)

    if i == 0:
        ax2.legend(
            frameon=False,
            fontsize=8,
            loc="lower right"
        )

fig.suptitle(
    "Ancestral 5-mer context shapes C→T substitution frequency",
    fontsize=15,
    fontweight="bold"
)

save(fig, "Figure4_5mer_landscape_REVISED")

# Build a revised five-slide presentation.
figure_names = [
    "Figure1_substitution_spectrum",
    "Figure2_posterior_sensitivity",
    "Figure3_chromosome_consistency",
    "Figure4_5mer_landscape_REVISED",
    "Figure5_motif_correlations"
]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

for name in figure_names:
    path = OUT / (name + ".png")

    if not path.exists():
        raise FileNotFoundError(path)

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    with Image.open(path) as img:
        width_px, height_px = img.size

    max_width = 12.5
    max_height = 6.8

    scale = min(
        max_width / width_px,
        max_height / height_px
    )

    width = width_px * scale
    height = height_px * scale

    slide.shapes.add_picture(
        str(path),
        Inches((13.333 - width) / 2),
        Inches((7.5 - height) / 2),
        width=Inches(width),
        height=Inches(height)
    )

ppt_path = OUT / "Phase1_Figures_REVISED.pptx"
prs.save(ppt_path)

print("CREATED:", ppt_path)
print("FIGURE REFINEMENT: PASS")
