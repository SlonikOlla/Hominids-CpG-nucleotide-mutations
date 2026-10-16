#!/usr/bin/env python3

import csv
from pathlib import Path

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from pptx import Presentation
from pptx.util import Inches

ROOT = Path("results/phase1_genome/aggregated")
OUT = ROOT / "figures"
OUT.mkdir(parents=True, exist_ok=True)

LINEAGES = ["human", "chimp", "bonobo", "gorilla"]
LABELS = ["Human", "Chimpanzee", "Bonobo", "Gorilla"]
CONTEXTS = ["CG", "CHG", "CHH"]

COLORS = {
    "CG": "#B44E49",
    "CHG": "#427EAD",
    "CHH": "#549B73"
}

LINEAGE_COLORS = {
    "human": "#B44E49",
    "chimp": "#427EAD",
    "bonobo": "#549B73",
    "gorilla": "#8A67A5"
}

plt.rcParams.update({
    "font.size": 10,
    "axes.spines.top": False,
    "axes.spines.right": False,
    "pdf.fonttype": 42,
    "ps.fonttype": 42,
    "figure.facecolor": "white",
    "savefig.facecolor": "white"
})

def read_table(filename):
    path = ROOT / filename
    if not path.exists():
        raise FileNotFoundError(path)

    with open(path, newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))

def save_figure(fig, filename):
    fig.savefig(
        OUT / f"{filename}.png",
        dpi=300,
        bbox_inches="tight"
    )
    fig.savefig(
        OUT / f"{filename}.pdf",
        bbox_inches="tight"
    )
    plt.close(fig)
    print("CREATED:", filename)

contexts = read_table("genome_context_summary.tsv")
ratios = read_table("genome_context_ratios.tsv")
chromosomes = read_table("chromosome_CT_ratios.tsv")
motifs = read_table("genome_5mer_validation.tsv")
correlations = read_table("motif_lineage_correlations.tsv")

# ============================================================
# FIGURE 1: Genome-wide substitution spectrum
# ============================================================

fig, axes = plt.subplots(1, 3, figsize=(14, 4.5))

for ax, mutation in zip(axes, ["C>T", "C>A", "C>G"]):
    x = np.arange(4)
    width = 0.25

    for j, context in enumerate(CONTEXTS):
        values = []

        for lineage in LINEAGES:
            row = next(
                r for r in contexts
                if r["lineage"] == lineage
                and r["posterior_threshold"] == "0.95"
                and r["context"] == context
            )

            rate = (
                int(row[mutation]) /
                int(row["opportunities"]) * 1_000_000
            )
            values.append(rate)

        ax.bar(
            x + (j - 1) * width,
            values,
            width,
            color=COLORS[context],
            label=context
        )

    ax.set_xticks(x)
    ax.set_xticklabels(LABELS, rotation=30, ha="right")
    ax.set_title(mutation, fontsize=13, fontweight="bold")
    ax.set_ylabel("Substitutions per million opportunities")
    ax.grid(axis="y", alpha=0.2)
    ax.set_axisbelow(True)

axes[0].legend(frameon=False)
fig.suptitle(
    "Genome-wide substitution spectrum",
    fontsize=15,
    fontweight="bold"
)
fig.tight_layout()
save_figure(fig, "Figure1_substitution_spectrum")

# ============================================================
# FIGURE 2: Posterior threshold sensitivity
# ============================================================

fig, axes = plt.subplots(1, 2, figsize=(12, 4.5))

for ax, comparison in zip(axes, ["CG/CHG", "CG/CHH"]):

    for lineage in LINEAGES:
        values = []

        for threshold in ["0.90", "0.95", "0.99"]:
            row = next(
                r for r in ratios
                if r["lineage"] == lineage
                and r["posterior_threshold"] == threshold
                and r["comparison"] == comparison
            )
            values.append(float(row["CT_rate_ratio"]))

        ax.plot(
            ["0.90", "0.95", "0.99"],
            values,
            marker="o",
            linewidth=2,
            color=LINEAGE_COLORS[lineage],
            label=lineage.capitalize()
        )

    ax.set_title(comparison, fontsize=13)
    ax.set_xlabel("Minimum ancestral posterior probability")
    ax.set_ylabel("C→T substitution-rate ratio")
    ax.grid(alpha=0.2)

axes[0].legend(frameon=False)
fig.tight_layout()
save_figure(fig, "Figure2_posterior_sensitivity")

# ============================================================
# FIGURE 3: Chromosome-level consistency
# ============================================================

fig, axes = plt.subplots(1, 2, figsize=(14, 4.8))

for ax, comparison in zip(axes, ["CG/CHG", "CG/CHH"]):

    for lineage in LINEAGES:
        values = []

        for chromosome in range(1, 23):
            row = next(
                r for r in chromosomes
                if r["chromosome"] == f"chr{chromosome}"
                and r["lineage"] == lineage
                and r["comparison"] == comparison
            )
            values.append(float(row["CT_rate_ratio"]))

        ax.plot(
            range(1, 23),
            values,
            marker="o",
            markersize=3,
            linewidth=1.5,
            color=LINEAGE_COLORS[lineage],
            label=lineage.capitalize()
        )

    ax.set_title(comparison)
    ax.set_xlabel("Autosome")
    ax.set_ylabel("C→T substitution-rate ratio")
    ax.set_xticks(range(1, 23, 2))
    ax.grid(alpha=0.2)

axes[0].legend(frameon=False)
fig.tight_layout()
save_figure(fig, "Figure3_chromosome_consistency")

# ============================================================
# FIGURE 4: Ancestral 5-mer substitution landscape
# ============================================================

fig, axes = plt.subplots(
    3, 1, figsize=(15, 9),
    gridspec_kw={"height_ratios": [1, 1, 1]}
)

for ax, context in zip(axes, CONTEXTS):
    subset = [
        r for r in motifs if r["context"] == context
    ]

    motif_names = sorted(set(
        r["ancestral_5mer"] for r in subset
    ))

    matrix = np.array([
        [
            float(next(
                r["CT_rate_per_million"]
                for r in subset
                if r["lineage"] == lineage
                and r["ancestral_5mer"] == motif
            ))
            for motif in motif_names
        ]
        for lineage in LINEAGES
    ])

    normalized = matrix / matrix.mean(axis=1, keepdims=True)

    im = ax.imshow(
        normalized,
        aspect="auto",
        cmap="viridis",
        vmin=0.4,
        vmax=1.8
    )

    ax.set_yticks(range(4))
    ax.set_yticklabels(LABELS)
    ax.set_xticks([])
    ax.set_title(
        f"{context} — {len(motif_names)} ancestral 5-mers",
        fontsize=12
    )

    fig.colorbar(
        im, ax=ax,
        fraction=0.015,
        pad=0.01,
        label="Relative C→T rate"
    )

fig.suptitle(
    "Sequence-context dependence of C→T substitutions",
    fontsize=15,
    fontweight="bold"
)
fig.tight_layout()
save_figure(fig, "Figure4_5mer_landscape")

# ============================================================
# FIGURE 5: Conservation of motif rankings
# ============================================================

pairs = [
    ("human", "chimp"),
    ("human", "bonobo"),
    ("human", "gorilla"),
    ("chimp", "bonobo"),
    ("chimp", "gorilla"),
    ("bonobo", "gorilla")
]

matrix = np.array([
    [
        float(next(
            r["spearman_rho"]
            for r in correlations
            if r["context"] == context
            and r["lineage_1"] == a
            and r["lineage_2"] == b
        ))
        for a, b in pairs
    ]
    for context in CONTEXTS
])

fig, ax = plt.subplots(figsize=(11, 4.2))

im = ax.imshow(
    matrix,
    cmap="Blues",
    vmin=0.90,
    vmax=1.00,
    aspect="auto"
)

ax.set_xticks(range(len(pairs)))
ax.set_xticklabels(
    [f"{a}–{b}" for a, b in pairs],
    rotation=30,
    ha="right"
)
ax.set_yticks(range(3))
ax.set_yticklabels(CONTEXTS)

for i in range(matrix.shape[0]):
    for j in range(matrix.shape[1]):
        value = matrix[i, j]
        ax.text(
            j, i, f"{value:.3f}",
            ha="center",
            va="center",
            color="white" if value > 0.975 else "black",
            fontsize=10
        )

ax.set_title(
    "Cross-lineage conservation of ancestral 5-mer substitution rates",
    fontsize=13
)

fig.colorbar(im, ax=ax, label="Spearman correlation (ρ)")
fig.tight_layout()
save_figure(fig, "Figure5_motif_correlations")

# ============================================================
# POWERPOINT ASSEMBLY
# ============================================================

figure_names = [
    "Figure1_substitution_spectrum",
    "Figure2_posterior_sensitivity",
    "Figure3_chromosome_consistency",
    "Figure4_5mer_landscape",
    "Figure5_motif_correlations"
]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

from PIL import Image

for name in figure_names:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    path = OUT / f"{name}.png"

    with Image.open(path) as img:
        width_px, height_px = img.size

    max_width = 12.5
    max_height = 6.8
    scale = min(
        max_width / width_px,
        max_height / height_px
    )
    w = width_px * scale
    h = height_px * scale

    slide.shapes.add_picture(
        str(path),
        Inches((13.333 - w) / 2),
        Inches((7.5 - h) / 2),
        width=Inches(w),
        height=Inches(h)
    )

ppt_path = OUT / "Phase1_Figures.pptx"
prs.save(ppt_path)

print("CREATED:", ppt_path)
print("FIGURE GENERATION: PASS")
